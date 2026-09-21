"""本文件对外提供 DOCX 锚点段落删除工具：精确定位、原子保存并返回变更证据。

对外提供:
    observe_docx_delete_candidate — 观察锚点对应段落并生成仅本次运行可用的删除候选
    delete_docx_paragraph — 凭候选标识删除该段落并返回变更证据

输入:
    两个工具都声明 `runtime: ToolRuntime` 参数；受治理的空间上下文（workspace /
    content_ref / y）与本次运行的变更证据经 `runtime.context` 提供。

输出:
    观察返回候选标识、段落序号与段落文本；删除返回变更前后的哈希与段落序号。

具体工作流:
    (1) 由受治理上下文取载体引用与锚点纵坐标，并校验 write 权限
    (2) 载体归属判定委托 focus.security（经 ObservationService 解析），本模块不自行比较工作根
    (3) 复刻 focus.readers 的段落顺序定位目标段落，要求文本在文档中唯一
    (4) 原子替换：同目录临时文件写好后替换，并复核删除后的段落序列
    (5) 确定性拒绝记入本次运行的变更证据，失败以可修正的 ToolException 收口

效果声明：观察者只读载体，删除者读并写同一份载体；两者都声明为可结构化枚举，目标来自受治理
上下文而非工具参数。观察者不得声明写效果——过度声明会为并不存在的写操作请求批准。

示例:
    observed = observe_docx_delete_candidate.func(runtime=runtime)
    delete_docx_paragraph.func(candidate_id=observed["candidate_id"], runtime=runtime)
"""

from __future__ import annotations

import hashlib
import os
import tempfile
from pathlib import Path
from typing import Any

from collections.abc import Mapping

from docx import Document
from langchain.tools import ToolRuntime
from langchain_core.tools import ToolException, tool

from focus.security import canonical_target
from focus.security.effects import ResolvedFsEffect, declare_all_effects, structured_fs
from focus.security.governed import declare_governed_keys
from plugins.spatial_patrol.spatial import ObservationService
from plugins.spatial_patrol.spatial_context import require_carrier_field, require_spatial_value

# 变更证据与候选存储参与决策：前者判定"是否已产生可验证变更"，后者限定"本次运行可删哪些段落"
declare_governed_keys("docx_change_evidence", "docx_observation_candidates")


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _runtime_context(runtime: ToolRuntime[dict]) -> dict[str, Any]:
    context = getattr(runtime, "context", None)
    if not isinstance(context, dict):
        raise RuntimeError("缺少空间上下文: runtime.context 必须为 dict")
    missing = [key for key in ("workspace", "content_ref", "y") if key not in context]
    if missing:
        raise RuntimeError(f"缺少空间上下文字段: {', '.join(missing)}")
    return context


def _extracted_parts(document: Any) -> tuple[list[str], list[str]]:
    """复刻 focus.readers._read_docx 的段落、表格拼接顺序。"""
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables: list[str] = []
    for table in document.tables:
        for row in table.rows:
            tables.append("\t".join(cell.text for cell in row.cells))
        tables.append("")
    return paragraphs, tables


def _paragraph_index_at_y(document: Any, y: float) -> int:
    paragraphs, table_parts = _extracted_parts(document)
    all_parts = [*paragraphs, *table_parts]
    extracted = "\n".join(all_parts)
    if not extracted or not paragraphs:
        raise ToolException("DOCX 没有可删除的正文段落")

    normalized_y = max(0.0, min(1.0, float(y)))
    position = min(int(normalized_y * len(extracted)), len(extracted) - 1)
    offset = 0
    for index, text in enumerate(paragraphs):
        end = offset + len(text)
        if offset <= position < end:
            return index
        if position == end:
            raise ToolException("锚点位于段落边界，无法唯一确定目标段落")
        offset = end + 1

    raise ToolException("锚点落在不受支持的表格文本区")


def _remove_paragraph(paragraph: Any) -> None:
    """封闭 python-docx 私有 XML API：移除整个 ``<w:p>``。"""
    element = paragraph._element
    parent = element.getparent()
    if parent is None:
        raise RuntimeError("目标段落没有 XML 父节点")
    parent.remove(element)
    paragraph._p = paragraph._element = None


def _verify_saved_document(
    path: Path,
    before_paragraphs: list[str],
    paragraph_index: int,
) -> None:
    reopened = Document(str(path))
    expected = before_paragraphs[:paragraph_index] + before_paragraphs[paragraph_index + 1 :]
    actual = [paragraph.text for paragraph in reopened.paragraphs]
    if actual != expected:
        raise RuntimeError("临时 DOCX 复核失败：目标段落未被唯一删除或相邻段落发生变化")


def _failure_key(context: dict[str, Any], candidate_id: str) -> str:
    return f"{context['content_ref']}\0{float(context['y']):.17g}\0{candidate_id}"


def _record_rejection(context: dict[str, Any], candidate_id: str, error: ToolException) -> None:
    evidence = require_spatial_value(context, "docx_change_evidence")
    if not isinstance(evidence, dict):
        raise RuntimeError("受治理的 DOCX 变更证据必须是 dict")
    evidence.clear()
    evidence.update({
        "changed": False,
        "error": f"DOCX 修改未执行：{error}",
        "failure_key": _failure_key(context, candidate_id),
    })


def _candidate_id(file_hash: str, paragraph_index: int, text: str) -> str:
    payload = f"{file_hash}\0{paragraph_index}\0{text}".encode("utf-8")
    return _sha256(payload)


def _require_write_target(context: dict[str, Any]) -> Path:
    permissions = frozenset(context.get("permissions") or ["read"])
    if "write" not in permissions:
        raise PermissionError("当前运行未授权 write")

    content_ref = str(context["content_ref"])
    if Path(content_ref).suffix.lower() != ".docx":
        raise ToolException("当前仅支持 DOCX 结构化修改；.doc 及其他载体保持只读")
    target = ObservationService.resolve_path(str(context["workspace"]), content_ref, context)
    if not target.is_file():
        raise ToolException(f"DOCX 文件不存在: {content_ref}")
    return target


@tool
def observe_docx_delete_candidate(runtime: ToolRuntime[dict]) -> dict[str, Any]:
    """观察锚点对应的 DOCX 正文段落，并生成仅本次运行可用的删除候选标识。"""
    context = _runtime_context(runtime)
    run_evidence = require_spatial_value(context, "docx_change_evidence")
    failure_key = _failure_key(context, "observe")
    if isinstance(run_evidence, dict) and run_evidence.get("failure_key") == failure_key:
        raise ToolException(
            f"相同参数已被确定性拒绝，请勿重复调用；原因为：{run_evidence.get('error', '目标不可修改')}"
        )
    target = _require_write_target(context)
    try:
        before_bytes = target.read_bytes()
        file_hash = _sha256(before_bytes)
        document = Document(str(target))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        paragraph_index = _paragraph_index_at_y(document, float(context["y"]))
        text = paragraphs[paragraph_index]
        if not text:
            raise ToolException("锚点对应空段落，无法建立删除候选")
        if sum(paragraph == text for paragraph in paragraphs) != 1:
            raise ToolException("文档中存在多个相同段落，无法唯一确定目标，已拒绝修改")
    except ToolException as error:
        _record_rejection(context, "observe", error)
        raise

    candidate_id = _candidate_id(file_hash, paragraph_index, text)
    candidates = require_spatial_value(context, "docx_observation_candidates")
    if not isinstance(candidates, dict):
        raise RuntimeError("本次运行的 DOCX 候选存储必须是 dict")
    candidates[candidate_id] = {
        "content_ref": str(context["content_ref"]),
        "file_hash": file_hash,
        "paragraph_index": paragraph_index,
        "text": text,
    }
    return {
        "candidate_id": candidate_id,
        "paragraph_index": paragraph_index,
        "text": text,
    }


@tool
def delete_docx_paragraph(candidate_id: str, runtime: ToolRuntime[dict]) -> dict[str, Any]:
    """凭本次运行的候选标识删除空间锚点对应的整个 DOCX 正文段落。"""
    context = _runtime_context(runtime)
    run_evidence = context.get("docx_change_evidence")
    if not isinstance(candidate_id, str) or not candidate_id:
        raise ToolException("candidate_id 不能为空，请先观察目标段落")
    failure_key = _failure_key(context, candidate_id)
    if isinstance(run_evidence, dict) and run_evidence.get("failure_key") == failure_key:
        raise ToolException(
            f"相同参数已被确定性拒绝，请勿重复调用；原因为：{run_evidence.get('error', '目标不可修改')}"
        )
    target = _require_write_target(context)
    candidates = context.get("docx_observation_candidates")
    snapshot = candidates.get(candidate_id) if isinstance(candidates, dict) else None
    if not isinstance(snapshot, dict):
        error = ToolException("候选标识无效或已过期，请重新观察目标段落")
        _record_rejection(context, candidate_id, error)
        raise error

    try:
        before_bytes = target.read_bytes()
        before_hash = _sha256(before_bytes)
        if (
            snapshot.get("content_ref") != str(context["content_ref"])
            or snapshot.get("file_hash") != before_hash
        ):
            raise ToolException("锚点目标已变化，请重新观察目标段落")
        document = Document(str(target))
        before_paragraphs = [paragraph.text for paragraph in document.paragraphs]
        paragraph_index = _paragraph_index_at_y(document, float(context["y"]))
        candidate = before_paragraphs[paragraph_index]
        if (
            paragraph_index != snapshot.get("paragraph_index")
            or candidate != snapshot.get("text")
        ):
            raise ToolException("锚点目标已变化，请重新观察目标段落")
        if sum(text == candidate for text in before_paragraphs) != 1:
            raise ToolException("文档中存在多个相同段落，无法唯一确定目标，已拒绝修改")
    except ToolException as error:
        _record_rejection(context, candidate_id, error)
        raise

    _remove_paragraph(document.paragraphs[paragraph_index])
    descriptor, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}.", suffix=".docx", dir=str(target.parent)
    )
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        document.save(str(temp_path))
        _verify_saved_document(temp_path, before_paragraphs, paragraph_index)
        after_bytes = temp_path.read_bytes()
        after_hash = _sha256(after_bytes)
        if after_hash == before_hash:
            raise RuntimeError("临时 DOCX 与原文件哈希相同，未产生可验证变更")
        if _sha256(target.read_bytes()) != before_hash:
            error = ToolException("DOCX 在编辑期间已被其他程序修改，已拒绝覆盖")
            _record_rejection(context, candidate_id, error)
            raise error
        try:
            os.replace(temp_path, target)
        except PermissionError as error:
            raise PermissionError(
                "DOCX 正被其他程序占用，请关闭 WPS/Word 中打开的该文件后重试"
            ) from error
    finally:
        temp_path.unlink(missing_ok=True)

    evidence = {
        "changed": True,
        "paragraph_index": paragraph_index,
        "removed_text": candidate,
        "before_hash": before_hash,
        "after_hash": after_hash,
    }
    if isinstance(run_evidence, dict):
        run_evidence.clear()
        run_evidence.update(evidence)
    if isinstance(candidates, dict):
        candidates.pop(candidate_id, None)
    return evidence


def _recoverable_docx_error(error: ToolException) -> str:
    return f"DOCX 修改未执行：{error}"


observe_docx_delete_candidate.handle_tool_error = _recoverable_docx_error
delete_docx_paragraph.handle_tool_error = _recoverable_docx_error


def _docx_carrier(context: Mapping[str, Any]) -> Path:
    """领域解析：本次删除作用在哪一份 DOCX 上。"""
    workspace = str(require_carrier_field(context, "workspace"))
    content_ref = str(require_carrier_field(context, "content_ref"))
    return canonical_target(Path(workspace).resolve(), content_ref)


def _docx_read_targets(args: Mapping[str, Any], context: Mapping[str, Any]) -> ResolvedFsEffect:
    return ResolvedFsEffect(reads=(_docx_carrier(context),))


def _docx_edit_targets(args: Mapping[str, Any], context: Mapping[str, Any]) -> ResolvedFsEffect:
    target = _docx_carrier(context)
    return ResolvedFsEffect(reads=(target,), writes=(target,))


declare_all_effects([observe_docx_delete_candidate], structured_fs(_docx_read_targets))
declare_all_effects([delete_docx_paragraph], structured_fs(_docx_edit_targets))
