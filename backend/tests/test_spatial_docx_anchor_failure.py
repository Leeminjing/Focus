"""DOCX 字符锚点、可恢复拒绝证据与旧坐标保护回归。"""

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import HTTPException
from langchain.tools import ToolRuntime
from langchain_core.tools import ToolException
from pydantic import ValidationError

from plugins.spatial_patrol.docx_edit import (
    _paragraph_index_at_y,
    observe_docx_delete_candidate,
)
from plugins.spatial_patrol.routes import (
    TEXT_COORDINATE_SPACE,
    _anchor_payload,
    _spatial_run_error,
    _validate_write_anchor,
)
from focus.plugins.loader import _collect_assets
from focus.plugins.schemas import PluginManifest


def _save_document(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))


def _runtime(workspace: Path, y: float, evidence: dict) -> ToolRuntime:
    return ToolRuntime(
        state={},
        context={
            "workspace": str(workspace),
            "content_ref": "sample.docx",
            "page": 1,
            "x": 0.5,
            "y": y,
            "permissions": ["read", "write"],
            "run_id": "run-anchor-failure",
            "docx_change_evidence": evidence,
            "docx_observation_candidates": {},
        },
        config={}, stream_writer=None, tool_call_id=None, store=None, tools=[],
    )


def _real_failure_paragraphs() -> list[str]:
    target = "（含页眉层水印“Trae ai”）"
    total_length = 1321
    tail_length = total_length - 33 - len(target) - 2
    return ["首" * 33, target, "尾" * tail_length]


def test_anchor_models_are_complete_under_real_plugin_loader():
    plugin_dir = Path(__file__).parents[2] / "plugins" / "spatial-patrol"
    manifest = PluginManifest.model_validate(
        json.loads((plugin_dir / "plugin.json").read_text(encoding="utf-8"))
    )

    router = _collect_assets(plugin_dir, manifest)["router"]
    endpoint = next(route.endpoint for route in router.routes if route.name == "create_anchor")
    namespace = endpoint.__globals__
    model = namespace["AnchorCreate"]

    assert all(
        namespace[name].__pydantic_complete__ is True
        for name in ("AnchorCreate", "AnchorPatch", "CopyRequest")
    )
    body = model.model_validate({
        "task_id": "task-1",
        "content_ref": "sample.docx",
        "page": 1,
        "x": 0.5,
        "y": 0.1,
        "coordinate_space": TEXT_COORDINATE_SPACE,
    })
    assert body.coordinate_space == TEXT_COORDINATE_SPACE
    with pytest.raises(ValidationError):
        model.model_validate({
            "task_id": "task-1",
            "content_ref": "sample.docx",
            "x": 0.5,
            "y": 0.1,
            "coordinate_space": "pixel-ratio-v0",
        })


def test_real_pixel_ratio_hits_boundary_but_character_ratio_hits_target(tmp_path):
    paragraphs = _real_failure_paragraphs()
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    document = Document(str(target))
    extracted_length = len("\n".join(paragraphs))
    assert extracted_length == 1321

    with pytest.raises(ToolException, match="段落边界"):
        _paragraph_index_at_y(document, 0.024988)

    target_start = len(paragraphs[0]) + 1
    character_y = (target_start + len(paragraphs[1]) / 2) / extracted_length
    assert _paragraph_index_at_y(document, character_y) == 1

    true_boundary_y = (len(paragraphs[0]) + 0.25) / extracted_length
    with pytest.raises(ToolException, match="段落边界"):
        _paragraph_index_at_y(document, true_boundary_y)


def test_rejection_is_projected_and_identical_retry_is_short_circuited(tmp_path, monkeypatch):
    paragraphs = _real_failure_paragraphs()
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    evidence: dict = {}

    with pytest.raises(ToolException, match="段落边界"):
        observe_docx_delete_candidate.func(runtime=_runtime(tmp_path, 0.024988, evidence))

    assert evidence["changed"] is False
    assert "锚点位于段落边界" in evidence["error"]
    assert _spatial_run_error(None, evidence) == evidence["error"]
    anchor = SimpleNamespace(to_payload=lambda: {"spatial_id": "soldier-1"})
    run = SimpleNamespace(status="success", error=_spatial_run_error(None, evidence))
    assert _anchor_payload(anchor, run)["run_error"] == evidence["error"]

    def must_not_reopen(_path):
        raise AssertionError("相同确定性失败不应再次打开 DOCX")

    monkeypatch.setattr("plugins.spatial_patrol.docx_edit.Document", must_not_reopen)
    with pytest.raises(ToolException, match="相同参数已被确定性拒绝"):
        observe_docx_delete_candidate.func(runtime=_runtime(tmp_path, 0.024988, evidence))


def test_legacy_docx_anchor_cannot_receive_write_permission():
    legacy = SimpleNamespace(content_ref="sample.docx", region=None)
    with pytest.raises(HTTPException) as caught:
        _validate_write_anchor(legacy, ["read", "write"])
    assert caught.value.status_code == 409
    assert "重新点击目标文字" in caught.value.detail

    current = SimpleNamespace(
        content_ref="sample.docx",
        region={"coordinate_space": TEXT_COORDINATE_SPACE},
    )
    _validate_write_anchor(current, ["read", "write"])
    _validate_write_anchor(legacy, ["read"])
