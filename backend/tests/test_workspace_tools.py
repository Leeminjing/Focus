import asyncio
import os
from pathlib import Path

import pytest
from langchain.tools import ToolRuntime
from langchain_core.messages import ToolMessage

from focus.tools.builtins.workspace_tools import (
    WORKSPACE_TOOLS,
    list_files,
    read_file,
    select_workspace_tools,
    write_file,
)


def _runtime(workspace, permissions):
    return ToolRuntime(
        state={},
        context={"workspace": str(workspace), "permissions": permissions},
        config={},
        stream_writer=None,
        tool_call_id=None,
        store=None,
        tools=[],
    )


def test_select_workspace_tools_by_permission(tmp_path):
    assert [t.name for t in select_workspace_tools(["read"])] == ["read_file", "list_files"]
    assert [t.name for t in select_workspace_tools(["read", "write"])] == ["read_file", "list_files", "write_file"]
    full = {t.name for t in select_workspace_tools(["read", "write", "host_command"])}
    assert full == {"read_file", "list_files", "write_file", "bash", "powershell", "cmd", "sh"}
    # 默认只读
    assert [t.name for t in select_workspace_tools(None)] == ["read_file", "list_files"]


def test_read_file_within_workspace(tmp_path):
    target = tmp_path / "hello.txt"
    target.write_text("内容", encoding="utf-8")
    runtime = _runtime(tmp_path, ["read"])
    assert read_file.func(path="hello.txt", runtime=runtime) == "内容"
    assert read_file.func(path=str(target), runtime=runtime) == "内容"


def test_workspace_path_type_errors_are_recoverable_tool_messages(tmp_path):
    target = tmp_path / "hello.txt"
    target.write_text("内容", encoding="utf-8")
    runtime = _runtime(tmp_path, ["read"])

    read_result = asyncio.run(read_file.ainvoke({
        "type": "tool_call",
        "id": "read-directory",
        "name": "read_file",
        "args": {"path": ".", "runtime": runtime},
    }))
    assert isinstance(read_result, ToolMessage)
    assert read_result.status == "error"
    assert read_result.tool_call_id == "read-directory"
    assert "list_files" in read_result.content

    list_result = asyncio.run(list_files.ainvoke({
        "type": "tool_call",
        "id": "list-file",
        "name": "list_files",
        "args": {"path": "hello.txt", "runtime": runtime},
    }))
    assert isinstance(list_result, ToolMessage)
    assert list_result.status == "error"
    assert list_result.tool_call_id == "list-file"
    assert "read_file" in list_result.content


def test_read_file_denied_without_read(tmp_path):
    target = tmp_path / "a.txt"
    target.write_text("x", encoding="utf-8")
    with pytest.raises(PermissionError, match="read"):
        read_file.func(path="a.txt", runtime=_runtime(tmp_path, []))


def test_read_file_outside_workspace(tmp_path):
    outside = tmp_path.parent / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    with pytest.raises(PermissionError, match="不属于当前工作区"):
        read_file.func(path=str(outside), runtime=_runtime(tmp_path, ["read"]))


def test_write_file_permission_gate(tmp_path):
    runtime = _runtime(tmp_path, ["read"])
    with pytest.raises(PermissionError, match="write"):
        write_file.func(path="new.txt", content="x", runtime=runtime)

    ok = write_file.func(path="sub/new.txt", content="内容", runtime=_runtime(tmp_path, ["read", "write"]))
    assert "已写入" in ok
    assert (tmp_path / "sub" / "new.txt").read_text(encoding="utf-8") == "内容"


@pytest.mark.skipif(os.name != "nt", reason="Windows extended path prefix regression")
def test_write_file_accepts_equivalent_windows_extended_path(tmp_path, monkeypatch):
    target = tmp_path / "result" / "contract_markdown.md"
    real_resolve = Path.resolve

    def resolve_with_extended_target(self, *args, **kwargs):
        resolved = real_resolve(self, *args, **kwargs)
        if self == target:
            return Path("\\\\?\\" + str(resolved))
        return resolved

    monkeypatch.setattr(Path, "resolve", resolve_with_extended_target)
    write_file.func(
        path=str(target),
        content="合同",
        runtime=_runtime(tmp_path, ["read", "write"]),
    )
    assert target.read_text(encoding="utf-8") == "合同"


def test_missing_workspace_context():
    runtime = ToolRuntime(
        state={}, context={}, config={}, stream_writer=None,
        tool_call_id=None, store=None, tools=[],
    )
    with pytest.raises(RuntimeError, match="workspace"):
        read_file.func(path="a.txt", runtime=runtime)


def test_read_file_docx_dispatch(tmp_path):
    # 文档分发：.docx 经 focus/readers._read_docx 解析为可读文本
    from docx import Document

    doc = Document()
    doc.add_paragraph("段落内容")
    doc.add_table(rows=1, cols=2)
    doc.tables[0].rows[0].cells[0].text = "单元格A"
    doc.tables[0].rows[0].cells[1].text = "单元格B"
    target = tmp_path / "doc.docx"
    doc.save(str(target))

    result = read_file.func(path="doc.docx", runtime=_runtime(tmp_path, ["read"]))
    assert "段落内容" in result
    assert "单元格A" in result and "单元格B" in result


def test_read_file_unknown_extension_falls_back_to_utf8(tmp_path):
    target = tmp_path / "data.txt"
    target.write_text("普通文本", encoding="utf-8")
    result = read_file.func(path="data.txt", runtime=_runtime(tmp_path, ["read"]))
    assert result == "普通文本"


def test_all_tools_registered():
    assert [t.name for t in WORKSPACE_TOOLS] == [
        "read_file", "list_files", "write_file", "bash", "powershell", "cmd", "sh",
    ]
