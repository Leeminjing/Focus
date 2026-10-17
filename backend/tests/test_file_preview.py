"""File preview boundaries, immutable exports, and real DOCX/PDF extraction."""
import asyncio
from hashlib import sha256
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

from backend.app.desktop import file_routes as files


def test_file_boundaries_and_missing_files(tmp_path):
    nested = tmp_path / "reports" / "data.json"
    nested.parent.mkdir()
    nested.write_text('{"ok": true}', encoding="utf-8")
    assert files.resolve_file(str(tmp_path), "reports/data.json") == nested.resolve()
    with pytest.raises(HTTPException) as outside:
        files.resolve_file(str(tmp_path), "../outside.txt")
    assert outside.value.status_code == 403
    with pytest.raises(HTTPException) as missing:
        files.resolve_file(str(tmp_path), "missing.xlsx")
    assert missing.value.status_code == 404


def test_historical_exports_do_not_drift_or_bypass_redaction(tmp_path):
    path = tmp_path / "plan.txt"
    path.write_bytes(b"new disk content")
    snapshot = {"after": {"available": True, "text": "old version\nsecret=***", "redacted": True}}
    assert files.read_content(path, snapshot, False) == b"old version\nsecret=***"
    assert files.read_content(path, snapshot, True) == b"new disk content"
    path.unlink()
    assert files.read_content(None, snapshot, False).startswith(b"old version")


def test_binary_versions_require_matching_bytes_and_preview_is_bounded(tmp_path):
    path = tmp_path / "image.png"
    original = b"\0original image"
    path.write_bytes(original)
    version = {"sha256": sha256(original).hexdigest(), "after": {"available": False, "reason": "binary"}}
    assert files.read_content(path, version, False) == original
    path.write_bytes(b"\0replacement image")
    with pytest.raises(HTTPException) as changed:
        files.read_content(path, version, False)
    assert changed.value.status_code == 409
    with pytest.raises(HTTPException) as oversized:
        files.read_content(path, None, False, limit=4)
    assert oversized.value.status_code == 413


def test_artifact_read_requires_permission_before_loading_snapshot(monkeypatch):
    class Denied:
        async def resolve(self, *args):
            return SimpleNamespace(permissions=frozenset())
    monkeypatch.setattr(files, "LoopLiveAccessPolicy", Denied)
    with pytest.raises(HTTPException) as denied:
        asyncio.run(files.resolve_target(None, files.FileTarget(loop_id="loop", fact_id="fact", revision=1)))
    assert denied.value.status_code == 403


def test_retained_snapshot_survives_slot_removal_but_current_file_does_not(monkeypatch):
    class Allowed:
        async def resolve(self, *args):
            return SimpleNamespace(permissions=frozenset({"read"}))
    class Reader:
        async def read(self, *args):
            return {"slot_id": "gone", "path": "plan.md", "after": {"available": True, "text": "saved"}}
    class Session:
        async def get(self, model, key):
            return SimpleNamespace(workspace_id="workspace") if model is files.AgentLoop else None
    monkeypatch.setattr(files, "LoopLiveAccessPolicy", Allowed)
    monkeypatch.setattr(files, "FileVersionReader", Reader)
    target = files.FileTarget(loop_id="loop", fact_id="fact", revision=1)
    path, version = asyncio.run(files.resolve_target(Session(), target))
    assert path is None and files.read_content(path, version, False) == b"saved"
    with pytest.raises(HTTPException) as missing:
        asyncio.run(files.resolve_target(Session(), target.model_copy(update={"current": True})))
    assert missing.value.status_code == 409


def test_docx_text_fallback_includes_paragraphs_and_tables_once(monkeypatch):
    import io
    from docx import Document
    document = Document()
    document.add_paragraph("Read only fallback")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table content"
    stream = io.BytesIO()
    document.save(stream)
    async def content(*args):
        return "report.docx", stream.getvalue()
    monkeypatch.setattr(files, "preview_bytes", content)
    result = asyncio.run(files.file_text(files.FileTarget(task_id="t", content_ref="report.docx"), None))
    assert result["content"].count("Read only fallback") == 1
    assert result["content"].count("Table content") == 1


def test_pdf_preview_uses_actual_page_count_and_rejects_invalid_page(monkeypatch):
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument.new()
    document.new_page(100, 200)
    document.new_page(100, 200)
    import io
    stream = io.BytesIO()
    document.save(stream)
    document.close()
    async def content(*args):
        return "report.pdf", stream.getvalue()
    monkeypatch.setattr(files, "preview_bytes", content)
    target = files.FileTarget(task_id="t", content_ref="report.pdf")
    result = asyncio.run(files.file_pdf_page(target, None, page=2))
    assert result["page_count"] == 2
    assert result["image"].startswith("iVBOR")
    with pytest.raises(HTTPException) as invalid:
        asyncio.run(files.file_pdf_page(target, None, page=3))
    assert invalid.value.status_code == 422
