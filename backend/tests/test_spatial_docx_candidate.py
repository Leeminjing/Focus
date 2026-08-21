"""DOCX 删除候选快照协议回归测试。"""

import inspect
from pathlib import Path

import pytest
from docx import Document
from langchain.tools import ToolRuntime
from langchain_core.tools import ToolException

from plugins.spatial_patrol.docx_edit import (
    delete_docx_paragraph,
    observe_docx_delete_candidate,
)
from plugins.spatial_patrol.routes import _launch_spatial_run


def _runtime(workspace: Path, y: float, candidates: dict, evidence=None) -> ToolRuntime:
    return ToolRuntime(
        state={},
        context={
            "workspace": str(workspace),
            "content_ref": "sample.docx",
            "page": 1,
            "x": 0.5,
            "y": y,
            "permissions": ["read", "write"],
            "run_id": "run-candidate",
            "docx_change_evidence": evidence if evidence is not None else {},
            "docx_observation_candidates": candidates,
        },
        config={}, stream_writer=None, tool_call_id=None, store=None, tools=[],
    )


def _save_document(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))


def _paragraph_y(paragraphs: list[str], index: int) -> float:
    extracted = "\n".join(paragraphs)
    start = sum(len(text) + 1 for text in paragraphs[:index])
    return (start + len(paragraphs[index]) / 2) / len(extracted)


def test_agent_tools_expose_candidate_id_protocol():
    observe_schema = observe_docx_delete_candidate.tool_call_schema.model_json_schema()
    delete_schema = delete_docx_paragraph.tool_call_schema.model_json_schema()
    launch_source = inspect.getsource(_launch_spatial_run)

    assert observe_schema["properties"] == {}
    assert set(delete_schema["properties"]) == {"candidate_id"}
    assert "expected_text" not in delete_schema["properties"]
    assert "observe_docx_delete_candidate" in launch_source
    assert '"docx_observation_candidates"' in launch_source


def test_candidate_id_deletes_original_curly_quote_paragraph(tmp_path):
    paragraphs = ["小花仙与赛尔号：关系探讨", "（含页眉层水印“Trae ai”）", "摘要"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    candidates: dict = {}
    runtime = _runtime(tmp_path, _paragraph_y(paragraphs, 1), candidates)

    observation = observe_docx_delete_candidate.func(runtime=runtime)
    assert observation["text"] == paragraphs[1]
    assert '"Trae ai"' != observation["text"]

    result = delete_docx_paragraph.func(
        candidate_id=observation["candidate_id"], runtime=runtime
    )

    reopened = Document(str(target))
    assert [paragraph.text for paragraph in reopened.paragraphs] == [paragraphs[0], paragraphs[2]]
    assert result["changed"] is True
    assert result["removed_text"] == paragraphs[1]


def test_unknown_candidate_is_rejected_without_changing_file(tmp_path):
    paragraphs = ["保留", "目标", "尾段"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    before = target.read_bytes()
    runtime = _runtime(tmp_path, _paragraph_y(paragraphs, 1), {})

    with pytest.raises(ToolException, match="无效或已过期"):
        delete_docx_paragraph.func(candidate_id="missing", runtime=runtime)

    assert target.read_bytes() == before


def test_file_change_after_observation_is_rejected_without_overwrite(tmp_path):
    paragraphs = ["保留", "原始目标", "尾段"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    candidates: dict = {}
    runtime = _runtime(tmp_path, _paragraph_y(paragraphs, 1), candidates)
    observation = observe_docx_delete_candidate.func(runtime=runtime)

    changed = Document(str(target))
    changed.paragraphs[1].text = "外部修改后的目标"
    changed.save(str(target))
    externally_changed_bytes = target.read_bytes()

    with pytest.raises(ToolException, match="目标已变化"):
        delete_docx_paragraph.func(
            candidate_id=observation["candidate_id"], runtime=runtime
        )

    assert target.read_bytes() == externally_changed_bytes
