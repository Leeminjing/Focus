"""spatial-patrol DOCX 精确删段落与原子保存回归测试。"""

from pathlib import Path

import pytest
from docx import Document
from langchain.tools import ToolRuntime
from langchain_core.tools import ToolException

from plugins.spatial_patrol.docx_edit import (
    delete_docx_paragraph,
    observe_docx_delete_candidate,
)


def _runtime(
    workspace: Path,
    y: float,
    permissions: list[str],
    evidence=None,
    candidates=None,
) -> ToolRuntime:
    return ToolRuntime(
        state={},
        context={
            "workspace": str(workspace),
            "content_ref": "sample.docx",
            "page": 1,
            "x": 0.5,
            "y": y,
            "permissions": permissions,
            "run_id": "run-docx",
            "docx_change_evidence": evidence if evidence is not None else {},
            "docx_observation_candidates": candidates if candidates is not None else {},
        },
        config={}, stream_writer=None, tool_call_id=None, store=None, tools=[],
    )


def _save_document(path: Path, paragraphs: list[str], table_text: str | None = None) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_text is not None:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = table_text
    document.save(str(path))


def _paragraph_y(paragraphs: list[str], index: int, trailing_parts=None) -> float:
    parts = [*paragraphs, *(trailing_parts or [])]
    extracted = "\n".join(parts)
    start = sum(len(text) + 1 for text in paragraphs[:index])
    return (start + len(paragraphs[index]) / 2) / len(extracted)


def test_delete_anchor_paragraph_preserves_neighbors_and_returns_evidence(tmp_path):
    paragraphs = ["标题", "（含页眉层水印“Trae ai”）", "摘要"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    evidence = {}
    candidates = {}
    runtime = _runtime(
        tmp_path, _paragraph_y(paragraphs, 1), ["read", "write"], evidence, candidates
    )
    observation = observe_docx_delete_candidate.func(runtime=runtime)

    result = delete_docx_paragraph.func(
        candidate_id=observation["candidate_id"], runtime=runtime,
    )

    reopened = Document(str(target))
    assert [paragraph.text for paragraph in reopened.paragraphs] == ["标题", "摘要"]
    assert result == evidence
    assert result["changed"] is True
    assert result["paragraph_index"] == 1
    assert result["removed_text"] == paragraphs[1]
    assert len(result["before_hash"]) == 64
    assert len(result["after_hash"]) == 64
    assert result["before_hash"] != result["after_hash"]


def test_delete_requires_write_permission_and_leaves_original_bytes(tmp_path):
    paragraphs = ["保留", "目标"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    before = target.read_bytes()

    with pytest.raises(PermissionError, match="write"):
        delete_docx_paragraph.func(
            candidate_id="not-observed",
            runtime=_runtime(tmp_path, _paragraph_y(paragraphs, 1), ["read"]),
        )

    assert target.read_bytes() == before


def test_delete_rejects_file_change_after_observation_and_leaves_changed_bytes(tmp_path):
    paragraphs = ["保留", "文件中的真实文字", "尾段"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    candidates = {}
    runtime = _runtime(
        tmp_path, _paragraph_y(paragraphs, 1), ["read", "write"], candidates=candidates
    )
    observation = observe_docx_delete_candidate.func(runtime=runtime)
    changed = Document(str(target))
    changed.paragraphs[1].text = "外部修改后的文字"
    changed.save(str(target))
    changed_bytes = target.read_bytes()

    with pytest.raises(ToolException, match="目标已变化"):
        delete_docx_paragraph.func(
            candidate_id=observation["candidate_id"], runtime=runtime,
        )

    assert target.read_bytes() == changed_bytes


def test_delete_rejects_table_anchor_and_leaves_original_bytes(tmp_path):
    paragraphs = ["正文"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs, table_text="表格目标")
    before = target.read_bytes()
    extracted = "\n".join(["正文", "表格目标", ""])
    table_y = (len("正文\n") + len("表格目标") / 2) / len(extracted)

    with pytest.raises(ToolException, match="表格文本区"):
        observe_docx_delete_candidate.func(
            runtime=_runtime(tmp_path, table_y, ["read", "write"]),
        )

    assert target.read_bytes() == before


def test_verify_failure_does_not_replace_original(tmp_path, monkeypatch):
    paragraphs = ["首段", "目标", "末段"]
    target = tmp_path / "sample.docx"
    _save_document(target, paragraphs)
    before = target.read_bytes()

    def fail_verify(*_args, **_kwargs):
        raise RuntimeError("模拟复核失败")

    monkeypatch.setattr("plugins.spatial_patrol.docx_edit._verify_saved_document", fail_verify)
    candidates = {}
    runtime = _runtime(
        tmp_path, _paragraph_y(paragraphs, 1), ["read", "write"], candidates=candidates
    )
    observation = observe_docx_delete_candidate.func(runtime=runtime)
    with pytest.raises(RuntimeError, match="模拟复核失败"):
        delete_docx_paragraph.func(
            candidate_id=observation["candidate_id"], runtime=runtime,
        )

    assert target.read_bytes() == before
    assert list(tmp_path.glob(".sample.*.docx")) == []
