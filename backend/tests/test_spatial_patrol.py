"""spatial-patrol 插件测试:观察服务单元、载体解析、插件接入校验(视觉硬性要求)。"""

import json
import shutil
from pathlib import Path

import pytest

from focus.plugins.interfaces import builtin_catalog
from focus.plugins.registry import PluginRegistry

PLUGIN_SOURCE = Path(__file__).resolve().parents[2] / "plugins" / "spatial-patrol"


def _copy_plugin(tmp_path: Path) -> Path:
    root = tmp_path / "plugins"
    target = root / "spatial-patrol"
    shutil.copytree(PLUGIN_SOURCE, target)
    return root


def _fake_config(with_vision: bool):
    from focus.config.app_config import AppConfig

    models = [{
        "name": "deepseek-text-only-test", "display_name": "text-only",
        "use": "focus.models.deepseek:DeepSeekChatOpenAI",
        "model": "deepseek-text-only-test", "context_window": 131072,
        "api_key": "sk-test", "base_url": "https://api.deepseek.com",
        "default": True,
    }]
    if with_vision:
        models.append({
            "name": "qwen3.7-plus", "display_name": "vision",
            "use": "langchain_openai:ChatOpenAI",
            "model": "qwen3.7-plus", "context_window": 128000,
            "api_key": "sk-vision-test",
            "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
        })
    return AppConfig(models=models)


def _load(root: Path, monkeypatch, with_vision: bool) -> PluginRegistry:
    # 插件内部经 plugins.spatial_patrol.vision.get_app_config 读配置,测试注入假配置
    config = _fake_config(with_vision)
    monkeypatch.setattr(
        "plugins.spatial_patrol.vision.get_app_config",
        lambda _path="config.yaml": config,
    )
    monkeypatch.setattr(
        "plugins.spatial_patrol.spatial.get_app_config",
        lambda _path="config.yaml": config,
    )
    registry = PluginRegistry(builtin_catalog())
    from focus.plugins.loader import load_plugins

    load_plugins(registry, root)
    return registry


# === 观察服务单元 ===


def test_resolve_path_containment():
    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({}, None)
    workspace = r"C:\ws\project"
    assert str(service.resolve_path(workspace, "docs/a.png")) == str(Path(r"C:\ws\project\docs\a.png"))
    with pytest.raises(ValueError, match="不属于当前工作区"):
        service.resolve_path(workspace, "..\\outside.png")


def test_is_viewable_suffixes():
    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({}, None)
    assert service.is_viewable("a.pdf")
    assert service.is_viewable("b.PNG")
    assert service.is_viewable("c.jpg")
    assert service.is_viewable("d.docx")
    assert service.is_viewable("e.txt")
    assert service.is_viewable("f.md")
    assert not service.is_viewable("g.bin")


def test_pdf_text_around_empty_page(tmp_path):
    """空白 PDF 页:get_text_bounded 返回空,服务返回「没有文字内容」。"""
    import asyncio

    import pypdfium2 as pdfium

    pdf_path = tmp_path / "empty.pdf"
    with pdfium.PdfDocument.new() as document:
        document.new_page(width=600, height=800)
        document.save(str(pdf_path))

    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({}, None)

    async def run():
        text = await service.observe(str(tmp_path), "empty.pdf", 1, 0.5, 0.5, 0.1)
        assert "没有文字内容" in text

    asyncio.run(run())


def test_observe_unsupported_suffix(tmp_path):
    import asyncio

    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({}, None)
    (tmp_path / "note.txt").write_text("hi", encoding="utf-8")

    async def run():
        with pytest.raises(ValueError, match="不支持的载体类型"):
            await service.observe(str(tmp_path), "note.txt", 1, 0.5, 0.5, 0.1)

    asyncio.run(run())


def test_radius_clamped_to_max():
    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({"max_radius": 0.5}, None)
    assert service.radius_start == 0.05
    assert service.radius_growth == 2.0


def test_supports_spatial_types():
    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({}, None)
    assert service.is_viewable("a.docx")
    assert service.is_viewable("b.md")
    assert service.is_viewable("c.txt")
    assert service.supports_spatial("a.docx")
    assert service.supports_spatial("a.doc")
    assert not service.supports_spatial("b.md")  # md/txt 纯预览,无空间能力
    assert not service.supports_spatial("c.txt")
    assert service.supports_spatial("d.pdf")
    assert service.supports_spatial("e.png")


def test_observe_docx_text_window(tmp_path):
    """docx 文本载体观察:锚点字符位置前后窗口(半径扩窗)。"""
    import asyncio

    from docx import Document

    doc_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("段落A" * 100)
    document.add_paragraph("段落B" * 100)
    document.save(str(doc_path))

    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({"max_radius": 0.5}, None)

    async def run():
        text = await service.observe(str(tmp_path), "sample.docx", 1, 0.5, 0.5, 0.1)
        assert "附近文本" in text
        assert "锚点字符位置" in text
        # 半径扩大 → 窗口更大
        small = await service.observe(str(tmp_path), "sample.docx", 1, 0.5, 0.5, 0.05)
        large = await service.observe(str(tmp_path), "sample.docx", 1, 0.5, 0.5, 0.2)
        assert len(large) > len(small)

    asyncio.run(run())


def test_observe_md_not_spatial(tmp_path):
    """md 不提供文本空间观察。"""
    import asyncio

    (tmp_path / "note.md").write_text("# 标题\n正文", encoding="utf-8")

    from plugins.spatial_patrol.spatial import ObservationService

    service = ObservationService({}, None)

    async def run():
        with pytest.raises(ValueError, match="不支持的载体类型"):
            await service.observe(str(tmp_path), "note.md", 1, 0.5, 0.5, 0.1)

    asyncio.run(run())


def test_spatial_write_run_requires_verified_change_for_done():
    from plugins.spatial_patrol.routes import _spatial_terminal_status

    assert _spatial_terminal_status(
        "success", requires_verified_change=True, change_evidence={}
    ) == "needs_action"
    assert _spatial_terminal_status(
        "success", requires_verified_change=True, change_evidence={"changed": True}
    ) == "done"
    assert _spatial_terminal_status(
        "success", requires_verified_change=False, change_evidence=None
    ) == "done"
    assert _spatial_terminal_status(
        "error", requires_verified_change=True, change_evidence={"changed": True}
    ) == "needs_action"


# === 插件接入(视觉硬性要求) ===


def test_plugin_unavailable_without_vision_model(tmp_path, monkeypatch):
    root = _copy_plugin(tmp_path)
    registry = _load(root, monkeypatch, with_vision=False)
    records = {item["name"]: item for item in registry.list_plugins()}
    assert records["spatial-patrol"]["status"] == "unavailable"
    assert "纯文本模型且无视觉插件" in records["spatial-patrol"]["reason"]
    assert registry.tools() == []
    assert registry.active_assets() == {}


def _write_fake_vision_plugin(root: Path) -> None:
    """写入一个提供 service.vision 的最小插件（作为视觉能力提供方的测试替身）。"""
    plugin_dir = root / "fake-vision"
    plugin_dir.mkdir(parents=True, exist_ok=True)
    (plugin_dir / "plugin.json").write_text(
        json.dumps({
            "name": "fake-vision", "version": "0.1.0", "enabled": True,
            "provides": ["service.vision"], "entry": "plugin.py",
        }),
        encoding="utf-8",
    )
    (plugin_dir / "plugin.py").write_text(
        "from focus.plugins.schemas import PluginDeclaration\n"
        "\n"
        "class _Vision:\n"
        "    async def describe(self, data_url):\n"
        "        return 'fake description'\n"
        "\n"
        "def build_plugin(context):\n"
        "    return PluginDeclaration(services={'service.vision': _Vision()})\n",
        encoding="utf-8",
    )


def test_plugin_active_with_vision_plugin_only(tmp_path, monkeypatch):
    """无 vision 条目但存在视觉插件(service.vision 提供者)→ spatial-patrol Active。"""
    root = _copy_plugin(tmp_path)
    _write_fake_vision_plugin(root)
    registry = _load(root, monkeypatch, with_vision=False)
    records = {item["name"]: item for item in registry.list_plugins()}
    assert records["spatial-patrol"]["status"] == "active"
    assert records["fake-vision"]["status"] == "active"


def test_plugin_active_with_vision_model(tmp_path, monkeypatch):
    root = _copy_plugin(tmp_path)
    registry = _load(root, monkeypatch, with_vision=True)
    records = {item["name"]: item for item in registry.list_plugins()}
    assert records["spatial-patrol"]["status"] == "active"
    # 观察工具不全局注入(仅空间小兵 agent_factory 装配,避免主 Agent 无空间上下文误调)
    assert registry.tools() == []
    assets = registry.active_assets()
    assert "spatial-patrol" in assets
    assert assets["spatial-patrol"]["router"] is not None
    assert records["spatial-patrol"]["desktop_assets"] == [
        "docx-editor.css", "docx-editor.js", "entry.js", "style.css", "viewer.js",
    ]


def test_plugin_disabled_no_assets(tmp_path, monkeypatch):
    root = _copy_plugin(tmp_path)
    manifest = root / "spatial-patrol" / "plugin.json"
    raw = json.loads(manifest.read_text(encoding="utf-8"))
    raw["enabled"] = False
    manifest.write_text(json.dumps(raw), encoding="utf-8")
    registry = _load(root, monkeypatch, with_vision=True)
    assert registry.active_assets() == {}
    assert registry.tools() == []
